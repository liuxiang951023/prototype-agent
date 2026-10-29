"""生成示例 docx 需求文档，便于快速测试 Agent。

用法：python scripts/make_sample_docx.py
会在 samples/ 目录下生成 sample_requirement.docx
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "samples"
OUTPUT_FILE = OUTPUT_DIR / "sample_requirement.docx"

SAMPLE_TEXT = """1. 系统概述
本系统是面向中小企业的客户关系管理（CRM）平台，支持客户档案管理、商机跟进、合同回款追踪等核心流程。

2. 功能模块

2.1 工作台首页
- 顶部统计卡片：本月新增客户数、待跟进商机数、待审批事项数、本月回款金额
- 待办事项列表：展示当前用户的待办任务，支持"处理"按钮
- 销售漏斗图：按阶段展示商机数量和金额

2.2 客户档案管理
- 客户列表：分页展示所有客户，支持按客户名称、行业、级别筛选
- 客户详情：顶部基础信息卡片（名称、行业、级别、负责人、联系电话）
- 跟进记录：客户档案下展示该客户的所有沟通记录时间轴
- 关联商机：展示该客户的所有商机列表

2.3 商机跟进
- 商机列表：展示所有商机，支持按阶段、客户、金额筛选
- 商机详情：展示商机阶段、金额、负责人、预计成交日期
- 阶段切换：商机可在"线索→初访→方案→报价→谈判→成交"6 个阶段间切换
- 跟进记录：商机下记录每次跟进的时间、方式、内容

2.4 合同回款
- 合同台账：展示所有合同，字段包括合同号、客户、金额、签订日期、状态
- 应收计划：展示每份合同的应收计划表，含计划日期、应收金额、实收金额
- 到账核销：录入到账记录，关联应收计划，自动标记核销状态

2.5 审批中心
- 待我审批：列出当前用户待审批的事项，支持"通过""驳回"操作
- 我发起的：列出当前用户发起的审批流程及当前节点
- 已处理：已审批的历史记录

3. 字段约定
- 客户分级：A（重点）/ B（普通）/ C（一般）
- 商机阶段：线索 / 初访 / 方案 / 报价 / 谈判 / 成交
- 合同状态：执行中 / 已完成 / 已终止
"""


def build_sample():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # 默认正文字体
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    for line in SAMPLE_TEXT.split("\n"):
        line = line.rstrip()
        if not line:
            continue
        # 简化处理：以数字开头视为标题
        if line and line[0].isdigit() and "." in line.split(" ")[0]:
            # 标题
            if line.startswith("1.") or line.startswith("2.") or line.startswith("3."):
                doc.add_heading(line, level=1)
            else:
                doc.add_heading(line, level=2)
        elif line.startswith("-"):
            doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(line)

    doc.save(OUTPUT_FILE)
    print(f"已生成示例 docx：{OUTPUT_FILE}")
    print(f"路径：{OUTPUT_FILE.as_posix()}")


if __name__ == "__main__":
    build_sample()
